"""Excel report generator for project deliverables."""
from __future__ import annotations

import io
import logging
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models.schemas import Project, ProjectProduct, Product, SymbolPlacement, BOQEntry

logger = logging.getLogger(__name__)

HEADER_FILL = PatternFill(start_color="003366", end_color="003366", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
SUBHEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
SUBHEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
SECTION_FILL = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
SECTION_FONT = Font(bold=True, size=10)
TOTAL_FILL = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
TOTAL_FONT = Font(bold=True, size=11)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
CURRENCY_FORMAT = '#,##0.00'
WRAP_ALIGNMENT = Alignment(wrap_text=True, vertical="center")


class ReportGenerator:
    async def generate_project_report(
        self,
        project_id: str,
        db: AsyncSession,
        cable_data: list[dict] | None = None,
        accessory_data: list[dict] | None = None,
    ) -> bytes:
        project = await db.get(Project, project_id)
        if not project:
            raise ValueError(f"Project not found: {project_id}")

        result = await db.execute(
            select(ProjectProduct)
            .where(ProjectProduct.project_id == project_id)
            .options(selectinload(ProjectProduct.product))
        )
        project_products = list(result.scalars().all())

        result = await db.execute(
            select(BOQEntry).where(BOQEntry.project_id == project_id)
        )
        boq_entries = list(result.scalars().all())

        result = await db.execute(
            select(SymbolPlacement)
            .join(SymbolPlacement.drawing)
            .where(SymbolPlacement.drawing.has(project_id=project_id))
        )
        placements = list(result.scalars().all())

        wb = openpyxl.Workbook()
        ws_cover = wb.active
        ws_cover.title = "Project Summary"
        self._create_cover_sheet(ws_cover, project, project_products)

        ws_products = wb.create_sheet("Product List")
        self._create_product_sheet(ws_products, project_products)

        ws_by_system = wb.create_sheet("By System Type")
        self._create_system_breakdown_sheet(ws_by_system, project_products)

        if cable_data:
            ws_cables = wb.create_sheet("Cable Schedule")
            self._create_cable_sheet(ws_cables, cable_data)

        if accessory_data:
            ws_accessories = wb.create_sheet("Accessories")
            self._create_accessory_sheet(ws_accessories, accessory_data)

        if boq_entries:
            ws_boq = wb.create_sheet("BOQ")
            self._create_boq_sheet(ws_boq, boq_entries)

        if placements:
            ws_placements = wb.create_sheet("Device Placements")
            self._create_placement_sheet(ws_placements, placements)

        ws_pricing = wb.create_sheet("Pricing Summary")
        self._create_pricing_sheet(ws_pricing, project_products, cable_data or [], accessory_data or [])

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.read()

    def _style_header_row(self, ws: openpyxl.worksheet.worksheet.Worksheet, row: int, cols: int) -> None:
        for col in range(1, cols + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")

    def _style_data_cell(self, ws: openpyxl.worksheet.worksheet.Worksheet, row: int, col: int) -> None:
        cell = ws.cell(row=row, column=col)
        cell.border = THIN_BORDER
        cell.alignment = WRAP_ALIGNMENT

    def _create_cover_sheet(self, ws, project: Project, products: list[ProjectProduct]) -> None:
        ws.column_dimensions["A"].width = 25
        ws.column_dimensions["B"].width = 45

        ws.merge_cells("A1:B1")
        title_cell = ws["A1"]
        title_cell.value = "PROJECT REPORT"
        title_cell.font = Font(bold=True, size=16, color="003366")
        title_cell.alignment = Alignment(horizontal="center")

        ws.merge_cells("A2:B2")
        subtitle = ws["A2"]
        subtitle.value = "Security & Building Systems Design"
        subtitle.font = Font(size=12, color="666666")
        subtitle.alignment = Alignment(horizontal="center")

        info_rows = [
            ("Project Name:", project.name),
            ("Client:", project.client_name or "N/A"),
            ("Location:", project.location or "N/A"),
            ("Status:", project.status or "Draft"),
            ("Created:", project.created_at.strftime("%Y-%m-%d") if project.created_at else "N/A"),
            ("Report Date:", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ]

        for i, (label, value) in enumerate(info_rows):
            row = i + 4
            ws.cell(row=row, column=1, value=label).font = Font(bold=True)
            ws.cell(row=row, column=2, value=value)
            ws.cell(row=row, column=1).border = THIN_BORDER
            ws.cell(row=row, column=2).border = THIN_BORDER

        summary_row = len(info_rows) + 5
        ws.cell(row=summary_row, column=1, value="SYSTEM SUMMARY").font = Font(bold=True, size=12)
        ws.cell(row=summary_row, column=1).fill = SECTION_FILL

        system_counts: dict[str, int] = {}
        total_cost = 0.0
        for pp in products:
            st = pp.system_type
            system_counts[st] = system_counts.get(st, 0) + pp.quantity
            if pp.product:
                total_cost += pp.product.price * pp.quantity

        for i, (system_type, count) in enumerate(system_counts.items()):
            row = summary_row + 1 + i
            ws.cell(row=row, column=1, value=system_type.replace("_", " ").title())
            ws.cell(row=row, column=2, value=f"{count} devices")
            ws.cell(row=row, column=1).border = THIN_BORDER
            ws.cell(row=row, column=2).border = THIN_BORDER

        total_row = summary_row + len(system_counts) + 2
        ws.cell(row=total_row, column=1, value="Total Products:").font = TOTAL_FONT
        ws.cell(row=total_row, column=2, value=sum(system_counts.values())).font = TOTAL_FONT
        ws.cell(row=total_row + 1, column=1, value="Estimated Cost:").font = TOTAL_FONT
        cost_cell = ws.cell(row=total_row + 1, column=2, value=total_cost)
        cost_cell.font = TOTAL_FONT
        cost_cell.number_format = CURRENCY_FORMAT

    def _create_product_sheet(self, ws, products: list[ProjectProduct]) -> None:
        headers = ["#", "Model Number", "Product Name", "Category", "System Type", "Qty", "Location", "Unit Price", "Total Price", "Cable Type", "Cable Length (m)", "Notes"]
        col_widths = [5, 20, 35, 18, 18, 6, 20, 12, 12, 18, 14, 25]

        for i, (header, width) in enumerate(zip(headers, col_widths)):
            ws.cell(row=1, column=i + 1, value=header)
            ws.column_dimensions[get_column_letter(i + 1)].width = width
        self._style_header_row(ws, 1, len(headers))

        for idx, pp in enumerate(products):
            row = idx + 2
            product = pp.product
            ws.cell(row=row, column=1, value=idx + 1)
            ws.cell(row=row, column=2, value=product.model_number if product else "")
            ws.cell(row=row, column=3, value=product.name if product else "")
            ws.cell(row=row, column=4, value=product.category if product else "")
            ws.cell(row=row, column=5, value=pp.system_type.replace("_", " ").title())
            ws.cell(row=row, column=6, value=pp.quantity)
            ws.cell(row=row, column=7, value=pp.location_note)
            price = product.price if product else 0
            ws.cell(row=row, column=8, value=price).number_format = CURRENCY_FORMAT
            ws.cell(row=row, column=9, value=price * pp.quantity).number_format = CURRENCY_FORMAT
            ws.cell(row=row, column=10, value=pp.cable_type)
            ws.cell(row=row, column=11, value=pp.cable_length_m)
            ws.cell(row=row, column=12, value=pp.notes)

            for col in range(1, len(headers) + 1):
                self._style_data_cell(ws, row, col)

        total_row = len(products) + 2
        ws.cell(row=total_row, column=5, value="TOTAL:").font = TOTAL_FONT
        ws.cell(row=total_row, column=6, value=sum(pp.quantity for pp in products)).font = TOTAL_FONT
        total_cost = sum((pp.product.price if pp.product else 0) * pp.quantity for pp in products)
        ws.cell(row=total_row, column=9, value=total_cost).font = TOTAL_FONT
        ws.cell(row=total_row, column=9).number_format = CURRENCY_FORMAT

    def _create_system_breakdown_sheet(self, ws, products: list[ProjectProduct]) -> None:
        by_system: dict[str, list[ProjectProduct]] = {}
        for pp in products:
            by_system.setdefault(pp.system_type, []).append(pp)

        headers = ["#", "Model Number", "Product Name", "Category", "Qty", "Unit Price", "Total"]
        col_widths = [5, 20, 35, 18, 6, 12, 12]
        for i, width in enumerate(col_widths):
            ws.column_dimensions[get_column_letter(i + 1)].width = width

        current_row = 1
        for system_type, items in by_system.items():
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=len(headers))
            section_cell = ws.cell(row=current_row, column=1, value=system_type.replace("_", " ").upper())
            section_cell.fill = SUBHEADER_FILL
            section_cell.font = SUBHEADER_FONT
            current_row += 1

            for i, header in enumerate(headers):
                ws.cell(row=current_row, column=i + 1, value=header)
            self._style_header_row(ws, current_row, len(headers))
            current_row += 1

            subtotal = 0.0
            for idx, pp in enumerate(items):
                product = pp.product
                ws.cell(row=current_row, column=1, value=idx + 1)
                ws.cell(row=current_row, column=2, value=product.model_number if product else "")
                ws.cell(row=current_row, column=3, value=product.name if product else "")
                ws.cell(row=current_row, column=4, value=product.category if product else "")
                ws.cell(row=current_row, column=5, value=pp.quantity)
                price = product.price if product else 0
                ws.cell(row=current_row, column=6, value=price).number_format = CURRENCY_FORMAT
                total = price * pp.quantity
                ws.cell(row=current_row, column=7, value=total).number_format = CURRENCY_FORMAT
                subtotal += total

                for col in range(1, len(headers) + 1):
                    self._style_data_cell(ws, current_row, col)
                current_row += 1

            ws.cell(row=current_row, column=5, value="Subtotal:").font = Font(bold=True)
            ws.cell(row=current_row, column=7, value=subtotal).font = Font(bold=True)
            ws.cell(row=current_row, column=7).number_format = CURRENCY_FORMAT
            current_row += 2

    def _create_cable_sheet(self, ws, cables: list[dict]) -> None:
        headers = ["#", "Cable Type", "Description", "Total Length (m)", "Roll Length (m)", "Qty Rolls", "System", "Purpose"]
        col_widths = [5, 22, 35, 15, 14, 10, 15, 30]

        for i, (header, width) in enumerate(zip(headers, col_widths)):
            ws.cell(row=1, column=i + 1, value=header)
            ws.column_dimensions[get_column_letter(i + 1)].width = width
        self._style_header_row(ws, 1, len(headers))

        from app.services.cable_calculator import CABLE_TYPES

        for idx, cable in enumerate(cables):
            row = idx + 2
            cable_type = cable.get("cable_type", "")
            cable_info = CABLE_TYPES.get(cable_type, {})
            ws.cell(row=row, column=1, value=idx + 1)
            ws.cell(row=row, column=2, value=cable_type)
            ws.cell(row=row, column=3, value=cable_info.get("description", cable_type))
            ws.cell(row=row, column=4, value=cable.get("total_length_m", 0))
            ws.cell(row=row, column=5, value=cable.get("roll_length_m", 305))
            ws.cell(row=row, column=6, value=cable.get("quantity_rolls", 1))
            ws.cell(row=row, column=7, value=cable.get("system_type", "").replace("_", " ").title())
            ws.cell(row=row, column=8, value=cable.get("purpose", ""))

            for col in range(1, len(headers) + 1):
                self._style_data_cell(ws, row, col)

    def _create_accessory_sheet(self, ws, accessories: list[dict]) -> None:
        headers = ["#", "Accessory Name", "Model Number", "Qty", "Purpose", "For Product", "System"]
        col_widths = [5, 30, 20, 6, 30, 20, 15]

        for i, (header, width) in enumerate(zip(headers, col_widths)):
            ws.cell(row=1, column=i + 1, value=header)
            ws.column_dimensions[get_column_letter(i + 1)].width = width
        self._style_header_row(ws, 1, len(headers))

        for idx, acc in enumerate(accessories):
            row = idx + 2
            ws.cell(row=row, column=1, value=idx + 1)
            ws.cell(row=row, column=2, value=acc.get("name", ""))
            ws.cell(row=row, column=3, value=acc.get("model_number", ""))
            ws.cell(row=row, column=4, value=acc.get("quantity", 1))
            ws.cell(row=row, column=5, value=acc.get("purpose", ""))
            ws.cell(row=row, column=6, value=acc.get("for_product", ""))
            ws.cell(row=row, column=7, value=acc.get("system_type", "").replace("_", " ").title())

            for col in range(1, len(headers) + 1):
                self._style_data_cell(ws, row, col)

    def _create_boq_sheet(self, ws, entries: list[BOQEntry]) -> None:
        headers = ["Item #", "Description", "Model Number", "Qty", "Unit", "Unit Price", "Total Price", "System Type", "Matched", "Source"]
        col_widths = [8, 35, 20, 6, 8, 12, 12, 15, 10, 20]

        for i, (header, width) in enumerate(zip(headers, col_widths)):
            ws.cell(row=1, column=i + 1, value=header)
            ws.column_dimensions[get_column_letter(i + 1)].width = width
        self._style_header_row(ws, 1, len(headers))

        for idx, entry in enumerate(entries):
            row = idx + 2
            ws.cell(row=row, column=1, value=entry.item_number)
            ws.cell(row=row, column=2, value=entry.description)
            ws.cell(row=row, column=3, value=entry.model_number)
            ws.cell(row=row, column=4, value=entry.quantity)
            ws.cell(row=row, column=5, value=entry.unit)
            ws.cell(row=row, column=6, value=entry.unit_price).number_format = CURRENCY_FORMAT
            ws.cell(row=row, column=7, value=entry.total_price).number_format = CURRENCY_FORMAT
            ws.cell(row=row, column=8, value=entry.system_type.replace("_", " ").title())
            ws.cell(row=row, column=9, value="Yes" if entry.matched_product_id else "No")
            ws.cell(row=row, column=10, value=entry.source_file)

            for col in range(1, len(headers) + 1):
                self._style_data_cell(ws, row, col)

    def _create_placement_sheet(self, ws, placements: list[SymbolPlacement]) -> None:
        headers = ["#", "Drawing", "System Type", "Symbol", "X", "Y", "Rotation", "Label", "Notes"]
        col_widths = [5, 20, 15, 15, 8, 8, 10, 20, 25]

        for i, (header, width) in enumerate(zip(headers, col_widths)):
            ws.cell(row=1, column=i + 1, value=header)
            ws.column_dimensions[get_column_letter(i + 1)].width = width
        self._style_header_row(ws, 1, len(headers))

        for idx, pl in enumerate(placements):
            row = idx + 2
            ws.cell(row=row, column=1, value=idx + 1)
            ws.cell(row=row, column=2, value=pl.drawing_id[:8])
            ws.cell(row=row, column=3, value=pl.system_type.replace("_", " ").title())
            ws.cell(row=row, column=4, value=pl.symbol_code)
            ws.cell(row=row, column=5, value=round(pl.x, 3))
            ws.cell(row=row, column=6, value=round(pl.y, 3))
            ws.cell(row=row, column=7, value=pl.rotation)
            ws.cell(row=row, column=8, value=pl.label)
            ws.cell(row=row, column=9, value=pl.notes)

            for col in range(1, len(headers) + 1):
                self._style_data_cell(ws, row, col)

    def _create_pricing_sheet(
        self,
        ws,
        products: list[ProjectProduct],
        cables: list[dict],
        accessories: list[dict],
    ) -> None:
        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 15

        ws.merge_cells("A1:B1")
        ws["A1"].value = "PRICING SUMMARY"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A1"].fill = HEADER_FILL
        ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")

        row = 3
        ws.cell(row=row, column=1, value="EQUIPMENT COSTS").font = SECTION_FONT
        ws.cell(row=row, column=1).fill = SECTION_FILL
        row += 1

        equipment_total = 0.0
        by_system: dict[str, float] = {}
        for pp in products:
            price = (pp.product.price if pp.product else 0) * pp.quantity
            equipment_total += price
            st = pp.system_type
            by_system[st] = by_system.get(st, 0) + price

        for system_type, cost in by_system.items():
            ws.cell(row=row, column=1, value=f"  {system_type.replace('_', ' ').title()}")
            ws.cell(row=row, column=2, value=cost).number_format = CURRENCY_FORMAT
            ws.cell(row=row, column=1).border = THIN_BORDER
            ws.cell(row=row, column=2).border = THIN_BORDER
            row += 1

        ws.cell(row=row, column=1, value="Equipment Subtotal:").font = Font(bold=True)
        ws.cell(row=row, column=2, value=equipment_total).font = Font(bold=True)
        ws.cell(row=row, column=2).number_format = CURRENCY_FORMAT
        row += 2

        from app.services.cable_calculator import CABLE_TYPES

        ws.cell(row=row, column=1, value="CABLE COSTS").font = SECTION_FONT
        ws.cell(row=row, column=1).fill = SECTION_FILL
        row += 1

        cable_total = 0.0
        for cable in cables:
            cable_type = cable.get("cable_type", "")
            cable_info = CABLE_TYPES.get(cable_type, {})
            qty = cable.get("quantity_rolls", 1)
            unit_price = cable_info.get("unit_price", 0)
            cost = qty * unit_price
            cable_total += cost
            ws.cell(row=row, column=1, value=f"  {cable_type}")
            ws.cell(row=row, column=2, value=cost).number_format = CURRENCY_FORMAT
            ws.cell(row=row, column=1).border = THIN_BORDER
            ws.cell(row=row, column=2).border = THIN_BORDER
            row += 1

        ws.cell(row=row, column=1, value="Cable Subtotal:").font = Font(bold=True)
        ws.cell(row=row, column=2, value=cable_total).font = Font(bold=True)
        ws.cell(row=row, column=2).number_format = CURRENCY_FORMAT
        row += 2

        ws.cell(row=row, column=1, value="GRAND TOTAL").font = TOTAL_FONT
        ws.cell(row=row, column=1).fill = TOTAL_FILL
        ws.cell(row=row, column=2, value=equipment_total + cable_total).font = TOTAL_FONT
        ws.cell(row=row, column=2).fill = TOTAL_FILL
        ws.cell(row=row, column=2).number_format = CURRENCY_FORMAT


report_generator = ReportGenerator()
