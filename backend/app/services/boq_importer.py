"""BOQ (Bill of Quantities) importer - supports Excel and CSV files."""
from __future__ import annotations

import csv
import io
import logging
import re
from pathlib import Path

import openpyxl
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.schemas import BOQEntry
from app.services.product_catalog import product_catalog

logger = logging.getLogger(__name__)

SYSTEM_TYPE_KEYWORDS = {
    "cctv": ["camera", "nvr", "dvr", "cctv", "surveillance", "dome", "bullet", "ptz", "ip cam", "recorder", "monitor"],
    "access_control": ["access", "card reader", "fingerprint", "face recognition", "controller", "electric lock", "turnstile", "biometric"],
    "gate": ["barrier", "gate", "bollard", "boom", "road blocker", "tire killer", "sliding gate"],
    "video_door_phone": ["intercom", "door phone", "door station", "indoor station", "video phone", "doorbell"],
    "fire": ["smoke", "fire", "heat detector", "alarm panel", "call point", "sounder", "fire alarm", "extinguisher"],
    "sound": ["speaker", "amplifier", "pa system", "horn", "audio", "sound", "microphone", "mixer"],
}


def detect_system_type(text: str) -> str:
    text_lower = text.lower()
    for system_type, keywords in SYSTEM_TYPE_KEYWORDS.items():
        for keyword in keywords:
            if keyword in text_lower:
                return system_type
    return "cctv"


def extract_model_number(text: str) -> str:
    patterns = [
        r"DS-[A-Z0-9\-/]+",
        r"iDS-[A-Z0-9\-/]+",
        r"[A-Z]{2,4}-[A-Z0-9\-/]+",
    ]
    for pattern in patterns:
        match = re.search(pattern, text, re.IGNORECASE)
        if match:
            return match.group(0)
    return ""


class BOQImporter:
    async def import_excel(
        self, file_content: bytes, filename: str, project_id: str, db: AsyncSession
    ) -> list[BOQEntry]:
        wb = openpyxl.load_workbook(io.BytesIO(file_content), read_only=True)
        entries = []

        for sheet in wb.sheetnames:
            ws = wb[sheet]
            headers = []
            header_row_found = False

            for row in ws.iter_rows(values_only=True):
                if not header_row_found:
                    row_values = [str(c).strip().lower() if c else "" for c in row]
                    if any(h in " ".join(row_values) for h in ["item", "description", "qty", "quantity", "model", "unit"]):
                        headers = row_values
                        header_row_found = True
                        continue
                    if not any(row):
                        continue
                    if len([c for c in row if c]) >= 3:
                        headers = [f"col_{i}" for i in range(len(row))]
                        header_row_found = True

                if header_row_found and any(row):
                    entry = self._parse_row(row, headers, project_id, filename)
                    if entry:
                        entries.append(entry)

        wb.close()

        for entry in entries:
            db.add(entry)
        await db.commit()

        return entries

    async def import_csv(
        self, file_content: bytes, filename: str, project_id: str, db: AsyncSession
    ) -> list[BOQEntry]:
        text = file_content.decode("utf-8", errors="replace")
        reader = csv.DictReader(io.StringIO(text))
        entries = []

        for i, row in enumerate(reader):
            description = row.get("description", row.get("Description", ""))
            model = row.get("model", row.get("Model", row.get("model_number", "")))
            qty_str = row.get("qty", row.get("quantity", row.get("Qty", row.get("Quantity", "1"))))
            unit = row.get("unit", row.get("Unit", "pcs"))
            price_str = row.get("unit_price", row.get("price", row.get("Unit Price", "0")))

            try:
                qty = int(float(qty_str)) if qty_str else 1
            except (ValueError, TypeError):
                qty = 1

            try:
                price = float(price_str) if price_str else 0.0
            except (ValueError, TypeError):
                price = 0.0

            if not description and not model:
                continue

            if not model:
                model = extract_model_number(description)

            entry = BOQEntry(
                project_id=project_id,
                item_number=i + 1,
                description=description or model,
                model_number=model,
                quantity=qty,
                unit=unit,
                unit_price=price,
                total_price=price * qty,
                system_type=detect_system_type(f"{description} {model}"),
                source_file=filename,
            )
            entries.append(entry)

        for entry in entries:
            db.add(entry)
        await db.commit()
        return entries

    def _parse_row(
        self, row: tuple, headers: list[str], project_id: str, filename: str
    ) -> BOQEntry | None:
        values = list(row)
        if len(values) < 2:
            return None

        data: dict = {}
        for i, header in enumerate(headers):
            if i < len(values) and values[i] is not None:
                val = str(values[i]).strip()
                if "item" in header or "no" in header or "num" in header:
                    try:
                        data["item_number"] = int(float(val))
                    except (ValueError, TypeError):
                        data["item_number"] = 0
                elif "desc" in header or "name" in header or "product" in header:
                    data["description"] = val
                elif "model" in header or "part" in header or "sku" in header:
                    data["model_number"] = val
                elif "qty" in header or "quantity" in header:
                    try:
                        data["quantity"] = int(float(val))
                    except (ValueError, TypeError):
                        data["quantity"] = 1
                elif "unit" in header and "price" not in header:
                    data["unit"] = val
                elif "price" in header or "cost" in header or "rate" in header:
                    try:
                        data["unit_price"] = float(val.replace(",", "").replace("$", ""))
                    except (ValueError, TypeError):
                        data["unit_price"] = 0.0
                elif "total" in header or "amount" in header:
                    try:
                        data["total_price"] = float(val.replace(",", "").replace("$", ""))
                    except (ValueError, TypeError):
                        pass

        description = data.get("description", "")
        model = data.get("model_number", "")

        if not description and not model:
            combined = " ".join(str(v) for v in values if v)
            if len(combined) < 5:
                return None
            description = combined

        if not model and description:
            model = extract_model_number(description)

        qty = data.get("quantity", 1)
        unit_price = data.get("unit_price", 0.0)

        return BOQEntry(
            project_id=project_id,
            item_number=data.get("item_number", 0),
            description=description or model,
            model_number=model,
            quantity=qty,
            unit=data.get("unit", "pcs"),
            unit_price=unit_price,
            total_price=data.get("total_price", unit_price * qty),
            system_type=detect_system_type(f"{description} {model}"),
            source_file=filename,
        )

    async def match_boq_to_products(
        self, entries: list[BOQEntry], db: AsyncSession
    ) -> list[BOQEntry]:
        for entry in entries:
            if entry.model_number:
                product = await product_catalog.get_product_by_model(db, entry.model_number)
                if product:
                    entry.matched_product_id = product.id
                    if not entry.system_type:
                        entry.system_type = product.system_type
        await db.commit()
        return entries


boq_importer = BOQImporter()
